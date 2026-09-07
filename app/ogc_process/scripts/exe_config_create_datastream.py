import base64
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font,  Alignment, Border, Side, NamedStyle, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import tempfile
import os


class ExeConfigCreateDatastream:

    @classmethod
    def is_valid_excel(cls, file_content):
        # Check the first bytes to determine whether it's an Excel file
        # .xlsx files usually start with the bytes 'PK'
        return file_content.startswith(b'PK')

    @classmethod
    def decode_base64_to_excel(cls, base64_string, output_file_path):
        file_content = base64.b64decode(base64_string)

        if not cls.is_valid_excel(file_content):
            raise ValueError(
                "The decoded content is not a valid Excel file.")

        # create a tmp file and test opening it
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as temp_file:
            temp_file.write(file_content)
            temp_file_path = temp_file.name
            load_workbook(temp_file_path)

        with open(output_file_path, "wb") as output_file:
            output_file.write(file_content)

    @classmethod
    def generateDatasreams(cls, file_xlsx):

        obsPConf = pd.read_excel(
            file_xlsx, sheet_name='1_observedProperty')
        sensorConf = pd.read_excel(file_xlsx, sheet_name='2_sensor')
        thingConf = pd.read_excel(file_xlsx, sheet_name='3_thing')

        dataStreamS = {"name": [], "Share": [], "nameShare": [], "description": [], "observationType": [
        ], "unit-name": [], "unit-symbol": [], "unit-definition": [], "nameUnique": []}

        Om_type = 'OM_Measurement'

        for idx, thing in thingConf.iterrows():
            thing = thing.dropna()

            for i in thing.index:
                if 'sensor' in i:
                    capteur = thing[i]
                    sensor = sensorConf[sensorConf.name ==
                                        capteur].T.dropna()

                    for j in sensor.index:
                        if 'observedProperty' in j:
                            obsP = sensor[sensor.index == j].values[0][0]

                            obsP_df = obsPConf[obsPConf.name == obsP]

                            dataStream = f"{thing['name']}_{capteur}_{obsP}"

                            dataStreamS['description'].append("")
                            dataStreamS['Share'].append("")
                            dataStreamS['nameShare'].append("")
                            dataStreamS['nameUnique'].append(dataStream)
                            dataStreamS['name'].append(dataStream)
                            dataStreamS['observationType'].append(Om_type)
                            dataStreamS['unit-name'].append(
                                obsP_df['unit-name'].values[0])
                            dataStreamS['unit-symbol'].append(
                                obsP_df['unit-symbol'].values[0])
                            dataStreamS['unit-definition'].append(
                                obsP_df['unit-definition'].values[0])

        dataStreamS = pd.DataFrame(dataStreamS)
        dataStreamS['observationProcedure'] = ''
        dataStreamS['graph'] = ''
        dataStreamS['frequency (min)'] = ''
        dataStreamS['minValue'] = ''
        dataStreamS['maxValue'] = ''

        col_to_move = dataStreamS.pop('nameUnique')
        dataStreamS['nameUnique'] = col_to_move

        cls.xmax = len(dataStreamS)+1

        with pd.ExcelWriter(file_xlsx, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            dataStreamS.to_excel(
                writer, index=False, sheet_name='4_datastream', freeze_panes=(1, 1))

    @staticmethod
    def get_or_create_named_style(wb, name):
        """
        openpyxl is finicky with styles....
        Retourne :
        - the existing NamedStyle object if present,
        - None if a style with this name exists only as a string (builtin),
        - a new NamedStyle object (already added to the wb) otherwise.
        """
        # homogeneous list of present names
        existing_names = [s.name if isinstance(
            s, NamedStyle) else s for s in wb.named_styles]

        # if a NamedStyle object exists -> return the object
        for s in wb.named_styles:
            if isinstance(s, NamedStyle) and s.name == name:
                return s

        # if the name exists only as a string (builtin style), we don't try to add it
        if name in existing_names:
            return None

        # otherwise create and add the NamedStyle
        st = NamedStyle(name=name)
        wb.add_named_style(st)
        return st

    @classmethod
    def stylage(cls, file_xlsx, config):

        xmin = 1
        xmax = cls.xmax
        workbook = load_workbook(filename=file_xlsx)
        sheet = workbook['4_datastream']
        sheet.freeze_panes = "B2"

        header_row = sheet[1]

        # get / create the style
        header = cls.get_or_create_named_style(workbook, "header")

        # desired attributes
        hdr_font = Font(bold=True)
        hdr_border = Border(bottom=Side(border_style="thin"))
        hdr_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)

        if header is None:
            # a style with that name already exists as a string -> apply directly to the cells
            for cell in header_row:
                cell.font = hdr_font
                cell.border = hdr_border
                cell.alignment = hdr_alignment
        else:
            # we have a reusable (or freshly created) NamedStyle: configure it then assign it
            header.font = hdr_font
            header.border = hdr_border
            header.alignment = hdr_alignment
            for cell in header_row:
                cell.style = header

        no_fill = PatternFill(bgColor="ff000000")
        dxf_no_fill = DifferentialStyle(fill=no_fill)
        final_rule = Rule(type="expression", dxf=dxf_no_fill, stopIfTrue=True)
        final_rule.formula = ['OR($B1="No",$B1="No but need Qualification")']
        sheet.conditional_formatting.add(f"C{xmin}:C{xmax}", final_rule)

        no_fill = PatternFill(bgColor="ffa4ffa4")
        dxf_no_fill = DifferentialStyle(fill=no_fill)
        final_rule = Rule(type="expression", dxf=dxf_no_fill, stopIfTrue=True)
        final_rule.formula = [
            'AND(OR($B1="No",OR($B1="No",$B1="No but need Qualification")), $D1<>"")']
        sheet.conditional_formatting.add(f"A{xmin}:D{xmax}", final_rule)

        no_fill = PatternFill(bgColor="ffa4ffa4")
        dxf_no_fill = DifferentialStyle(fill=no_fill)
        final_rule = Rule(type="expression", dxf=dxf_no_fill, stopIfTrue=True)
        final_rule.formula = [
            'AND($B2<>"",$C2<>"", $D2<>"")']
        sheet.conditional_formatting.add(f"A{xmin+1}:D{xmax}", final_rule)

        for col in ['B', 'C', 'D']:
            red_fill = PatternFill(bgColor="FFC7CE")
            dxf = DifferentialStyle(fill=red_fill)
            r = Rule(type="expression", dxf=dxf, stopIfTrue=False)
            r.formula = [f'${col}1=""']
            sheet.conditional_formatting.add(
                f"{col}{xmin}:{col}{xmax}", r)  # sheet.dimensions
            sheet.conditional_formatting.add(
                f"A{xmin}:A{xmax}", r)  # sheet.dimensions

        sheet.column_dimensions['A'].width = 40
        sheet.column_dimensions['E'].width = 20
        sheet.column_dimensions['N'].width = 40
        for col in ['C', 'D', 'H']:
            sheet.column_dimensions[col].width = 30
        for col in ['G', 'I']:
            sheet.column_dimensions[col].width = 15
        for col in ['B',  'J', 'K', 'L', 'M']:
            sheet.column_dimensions[col].width = 10

        alignment = Alignment(horizontal="center",
                              vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2, max_row=xmax+1):
            for cell in row:
                cell.alignment = alignment

        valeurs = config["AquisitionMethode"]

        validation = DataValidation(
            type="list", formula1=f'"{",".join(valeurs)}"')
        sheet.add_data_validation(validation)
        validation.add(f"I{xmin+1}:I{xmax}")

        valeurs = config["observationType"]
        validation2 = DataValidation(
            type="list", formula1=f'"{",".join(valeurs)}"')
        sheet.add_data_validation(validation2)
        validation2.add(f"E{xmin+1}:E{xmax}")

        valeurs = config["share"]
        validation2 = DataValidation(
            type="list", formula1=f'"{",".join(valeurs)}"')
        sheet.add_data_validation(validation2)
        validation2.add(f"B{xmin+1}:B{xmax}")

        valeurs = config['graph']
        validation2 = DataValidation(
            type="list", formula1=f'"{",".join(valeurs)}"')
        sheet.add_data_validation(validation2)
        validation2.add(f"J{xmin+1}:J{xmax}")

        workbook.active = workbook.sheetnames.index(sheet.title)
        workbook.save(file_xlsx)

    @classmethod
    def run(cls, base64_string, obs_name, config):

        # REMEMBER to strip trailing spaces if " " is last, delete

        date = datetime.now().strftime("%Y-%m-%dT%H:%M")
        output_file_path = f"app/static/xlsx/tmp/{obs_name}_config_{date}.xlsx"

        try:
            cls.decode_base64_to_excel(
                base64_string, output_file_path)
            cls.generateDatasreams(output_file_path)
            cls.stylage(output_file_path, config)

        except Exception as e:
            if str(e) == "1 must be greater than 2":
                e = "The file was not filled in"
            return {"error": str(e)}, 500

        return {
            "configFileUrl": f"static/xlsx/tmp/{obs_name}_config_{date}.xlsx"
        }, 200
